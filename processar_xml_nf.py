import os
import re
import xml.etree.ElementTree as ET
from pathlib import Path
from datetime import datetime

from openpyxl import load_workbook
from pypdf import PdfReader


# ============================================================
# CONFIGURAÇÃO
# ============================================================

PASTA_PROCESSAR = Path("A_PROCESSAR")
ARQUIVO_BANCO = Path("banco.xlsx")
ARQUIVO_RELATORIO = Path("relatorio_xml_pdf_v10.txt")


# ============================================================
# UTILITÁRIOS
# ============================================================

def limpar_cnpj(valor):
    if valor is None:
        return ""

    return re.sub(r"\D", "", str(valor))


def formatar_codigo(valor):
    if valor is None:
        return ""

    texto = str(valor).strip()

    if texto.endswith(".0"):
        texto = texto[:-2]

    return texto


def limpar_numero(valor):
    if valor is None:
        return None

    texto = str(valor).strip()
    texto = texto.replace(" ", "")

    numeros = re.sub(r"\D", "", texto)

    if not numeros:
        return None

    if len(numeros) > 12:
        return None

    return numeros.lstrip("0") or "0"


def nome_tag(tag):
    if not tag:
        return ""

    return tag.split("}")[-1].strip()


def normalizar_data(valor):
    if not valor:
        return None

    texto = str(valor).strip()

    formatos = [
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%dT%H:%M:%S.%f",
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d-%m-%Y",
    ]

    for formato in formatos:
        try:
            dt = datetime.strptime(texto[:26], formato)
            return dt.strftime("%d/%m/%Y")
        except Exception:
            pass

    m = re.search(
        r"(\d{2})[/-](\d{2})[/-](\d{4})",
        texto
    )

    if m:
        return (
            f"{m.group(1)}/"
            f"{m.group(2)}/"
            f"{m.group(3)}"
        )

    m = re.search(
        r"(\d{4})-(\d{2})-(\d{2})",
        texto
    )

    if m:
        return (
            f"{m.group(3)}/"
            f"{m.group(2)}/"
            f"{m.group(1)}"
        )

    return None


# ============================================================
# CHAVES DE ACESSO
# ============================================================

def encontrar_chaves(texto):
    """
    Procura chaves de acesso de 44 dígitos.

    Aceita:
    - chave contínua
    - chave separada por espaços
    """

    if not texto:
        return []

    texto = str(texto)

    encontrados = set()

    # Chave contínua
    for chave in re.findall(
        r"\b\d{44}\b",
        texto
    ):
        encontrados.add(chave)

    # Chave com espaços
    grupos = re.findall(
        r"(?:\d{4}\s*){11}",
        texto
    )

    for grupo in grupos:

        chave = re.sub(
            r"\D",
            "",
            grupo
        )

        if len(chave) == 44:
            encontrados.add(chave)

    return sorted(encontrados)


def validar_chave_nfe(chave):
    """
    Validação estrutural básica.

    A posição do modelo deve ser 55.
    """

    if not chave:
        return False

    if len(chave) != 44:
        return False

    if not chave.isdigit():
        return False

    modelo = chave[20:22]

    if modelo != "55":
        return False

    return True


def numero_por_chave_nfe(chave):
    """
    Extrai o número da NF-e da chave de acesso.

    Estrutura da chave:

    cUF
    AAMM
    CNPJ
    modelo
    série
    número
    etc.

    O número da NF-e ocupa as posições 25:34.
    """

    if not validar_chave_nfe(chave):
        return None

    numero = chave[25:34]

    if not numero.isdigit():
        return None

    numero = numero.lstrip("0")

    return numero or "0"


# ============================================================
# LEITURA DO XML
# ============================================================

def ler_xml(caminho):
    try:
        tree = ET.parse(caminho)
        root = tree.getroot()

        return root

    except Exception:
        return None


def todos_elementos(root):

    if root is None:
        return []

    return list(root.iter())


def extrair_textos_xml(root):

    dados = []

    for elem in todos_elementos(root):

        tag = nome_tag(elem.tag)
        texto = (elem.text or "").strip()

        if texto:

            dados.append({
                "tag": tag,
                "texto": texto
            })

        for chave, valor in elem.attrib.items():

            if valor:

                dados.append({
                    "tag": chave,
                    "texto": str(valor).strip()
                })

    return dados


# ============================================================
# IDENTIFICAÇÃO DO TIPO
# ============================================================

def identificar_tipo_xml(root):

    if root is None:
        return None

    tags = {
        nome_tag(elem.tag)
        for elem in todos_elementos(root)
    }

    # NF-e
    if (
        "infNFe" in tags
        or "NFe" in tags
    ):
        return "NF-e"

    # NFS-e
    if (
        "NFSe" in tags
        or "infNFSe" in tags
        or "nNFSe" in tags
        or "vISSQN" in tags
    ):
        return "NFS-e"

    # Procura modelo 55
    for elem in todos_elementos(root):

        tag = nome_tag(elem.tag)

        if tag.lower() == "mod":

            texto = (
                elem.text or ""
            ).strip()

            if texto == "55":
                return "NF-e"

    return None


# ============================================================
# NÚMERO DA NF — EXTRAÇÃO REDUNDANTE
# ============================================================

def candidatos_numero_xml(
    root,
    tipo,
    caminho_xml
):
    """V10: extração redundante, com fontes graduadas e sem inferência solta.

    Prioriza tags oficiais. Só usa aliases explícitos quando não existe número
    oficial; chaves de acesso continuam sendo uma segunda fonte forte para NF-e.
    Não usa números aleatórios do nome do arquivo.
    """
    candidatos = []
    dados = extrair_textos_xml(root)

    tags_oficiais = {
        "NF-e": {"nNF"},
        "NFS-e": {
            "nNFSe", "NumeroNfse", "NumeroNFSe",
            "numeroNfse", "numeroNFSe"
        },
    }
    tags_alternativas = {
        "NF-e": set(),
        "NFS-e": {
            "NumeroNota", "numeroNota", "NumeroNfseNacional",
            "numeroNfseNacional", "NumeroDocumento", "numeroDocumento",
            "Numero", "numero", "NumNota", "numNota", "nNota"
        },
    }

    def adicionar(tag_item, peso, sufixo=""):
        texto = str(tag_item.get("texto", "")).strip()
        if not re.fullmatch(r"\d{1,12}", texto):
            return
        numero = limpar_numero(texto)
        if not numero or numero == "0":
            return
        candidatos.append({
            "numero": numero,
            "metodo": f"tag:{tag_item['tag']}{sufixo}",
            "peso": peso
        })

    oficiais_encontrados = False
    for item in dados:
        if item["tag"] in tags_oficiais.get(tipo, set()):
            oficiais_encontrados = True
            adicionar(item, 100)

    # Aliases NFS-e só entram se não foi encontrada tag oficial.
    if not oficiais_encontrados:
        for item in dados:
            if item["tag"] in tags_alternativas.get(tipo, set()):
                adicionar(item, 75, "_alternativa")

    # Alguns provedores guardam o número em atributo de identificação.
    # Atributos só são usados quando ainda não há número candidato.
    if not candidatos and tipo == "NFS-e":
        aliases_attr = {"numero", "numeronfse", "numeronota", "idnfse"}
        for elem in root.iter():
            for nome, valor in elem.attrib.items():
                nome_limpo = nome_tag(nome)
                if nome_limpo.lower() not in aliases_attr:
                    continue
                item = {"tag": f"atributo:{nome_limpo}", "texto": str(valor).strip()}
                adicionar(item, 65, "_atributo")

    # Chaves podem estar no texto, em atributos ou no nome original do XML.
    texto_xml = " ".join(item.get("texto", "") for item in dados)
    atributos = " ".join(str(v) for elem in root.iter() for v in elem.attrib.values())
    chaves = encontrar_chaves(texto_xml + " " + atributos + " " + Path(caminho_xml).stem)
    chaves_validas = []
    for chave in chaves:
        if validar_chave_nfe(chave) and chave not in chaves_validas:
            chaves_validas.append(chave)

    if tipo == "NF-e":
        for chave in chaves_validas:
            numero = numero_por_chave_nfe(chave)
            if numero and numero != "0":
                candidatos.append({
                    "numero": numero,
                    "metodo": "chave_acesso",
                    "peso": 95,
                    "chave": chave
                })

    return candidatos

def decidir_numero_xml(
    root,
    tipo,
    caminho_xml
):

    candidatos = candidatos_numero_xml(
        root,
        tipo,
        caminho_xml
    )

    if not candidatos:

        return (
            None,
            [],
            "SEM_NÚMERO"
        )

    # --------------------------------------------------------
    # AGRUPAMENTO
    # --------------------------------------------------------

    agrupados = {}

    for candidato in candidatos:

        numero = candidato["numero"]

        if numero not in agrupados:
            agrupados[numero] = []

        agrupados[numero].append(
            candidato
        )

    # --------------------------------------------------------
    # RANKING
    # --------------------------------------------------------

    ranking = []

    for numero, itens in agrupados.items():

        peso_total = sum(
            item["peso"]
            for item in itens
        )

        ranking.append({

            "numero": numero,

            "peso": peso_total,

            "metodos": [
                item["metodo"]
                for item in itens
            ]
        })

    ranking.sort(
        key=lambda x: x["peso"],
        reverse=True
    )

    primeiro = ranking[0]

    # --------------------------------------------------------
    # CONFLITO
    # --------------------------------------------------------

    if len(ranking) > 1:

        if (
            ranking[0]["peso"]
            == ranking[1]["peso"]
        ):

            return (
                None,
                candidatos,
                "CONFLITO_NÚMEROS"
            )

    return (
        primeiro["numero"],
        candidatos,
        primeiro["metodos"]
    )


# ============================================================
# DATA DE EMISSÃO
# ============================================================

def extrair_data_xml(root):

    if root is None:
        return None

    candidatos = []

    tags_data = {
        "dhEmi",
        "dEmi",
        "dhEmissao",
        "dEmissao",
        "dataEmissao",
        "DataEmissao",
        "dtEmissao",
    }

    for elem in todos_elementos(root):

        tag = nome_tag(elem.tag)

        if tag in tags_data:

            texto = (
                elem.text or ""
            ).strip()

            data = normalizar_data(
                texto
            )

            if data:

                candidatos.append(
                    data
                )

    if candidatos:
        return candidatos[0]

    # Segunda tentativa
    for elem in todos_elementos(root):

        texto = (
            elem.text or ""
        ).strip()

        data = normalizar_data(
            texto
        )

        if data:
            return data

    return None


# ============================================================
# CNPJ DO EMITENTE / DESTINATÁRIO
# ============================================================

def extrair_emitente_destinatario(root):

    emit_cnpj = None
    dest_cnpj = None

    for elem in todos_elementos(root):

        tag = nome_tag(elem.tag)

        # ----------------------------------------------------
        # EMITENTE
        # ----------------------------------------------------

        if tag == "emit":

            for filho in elem.iter():

                if (
                    nome_tag(filho.tag)
                    == "CNPJ"
                ):

                    cnpj = limpar_cnpj(
                        filho.text
                    )

                    if len(cnpj) == 14:

                        emit_cnpj = cnpj

                        break

        # ----------------------------------------------------
        # DESTINATÁRIO
        # ----------------------------------------------------

        if tag == "dest":

            for filho in elem.iter():

                if (
                    nome_tag(filho.tag)
                    == "CNPJ"
                ):

                    cnpj = limpar_cnpj(
                        filho.text
                    )

                    if len(cnpj) == 14:

                        dest_cnpj = cnpj

                        break

    return (
        emit_cnpj,
        dest_cnpj
    )


# ============================================================
# CHAVE DO XML
# ============================================================

def extrair_chave_xml(
    root,
    caminho_xml
):

    textos = []

    if root is not None:

        for elem in todos_elementos(root):

            texto = (
                elem.text or ""
            ).strip()

            if texto:
                textos.append(texto)

            for valor in elem.attrib.values():

                if valor:
                    textos.append(
                        str(valor)
                    )

    textos.append(
        caminho_xml.name
    )

    chaves = encontrar_chaves(
        " ".join(textos)
    )

    # Prioriza modelo 55
    chaves_validas = [
        chave
        for chave in chaves
        if validar_chave_nfe(chave)
    ]

    if chaves_validas:
        return chaves_validas[0]

    if chaves:
        return chaves[0]

    return None


# ============================================================
# BANCO DE CLIENTES
# ============================================================

def carregar_clientes():

    wb = load_workbook(
        ARQUIVO_BANCO,
        data_only=True
    )

    ws = wb.active

    clientes = []

    headers = {}

    for coluna in range(
        1,
        ws.max_column + 1
    ):

        valor = ws.cell(
            row=1,
            column=coluna
        ).value

        if valor:

            headers[
                str(valor)
                .strip()
                .upper()
            ] = coluna

    col_cod = headers.get("COD")
    col_razao = headers.get(
        "RAZÃO SOCIAL"
    )
    col_cnpj = headers.get("CNPJ")
    col_caminho = headers.get(
        "CAMINHO ABSOLUTO"
    )

    for linha in range(
        2,
        ws.max_row + 1
    ):

        cod = formatar_codigo(
            ws.cell(
                linha,
                col_cod
            ).value
        )

        razao = str(
            ws.cell(
                linha,
                col_razao
            ).value or ""
        ).strip()

        cnpj = limpar_cnpj(
            ws.cell(
                linha,
                col_cnpj
            ).value
        )

        caminho = str(
            ws.cell(
                linha,
                col_caminho
            ).value or ""
        ).strip()

        if not cnpj:
            continue

        clientes.append({

            "cod": cod,

            "razao": razao,

            "cnpj": cnpj,

            "caminho": Path(caminho)
        })

    return clientes


# ============================================================
# IDENTIFICAÇÃO DO CLIENTE
# ============================================================

def identificar_cliente(
    clientes,
    emit_cnpj,
    dest_cnpj
):

    resultados = []

    # --------------------------------------------------------
    # DESTINATÁRIO = ENTRADA
    # --------------------------------------------------------

    if dest_cnpj:

        for cliente in clientes:

            if (
                cliente["cnpj"]
                == dest_cnpj
            ):

                resultados.append(
                    (
                        "ENTRADAS",
                        cliente
                    )
                )

    if len(resultados) == 1:
        return resultados[0]

    if len(resultados) > 1:

        return (
            "AMBIGUO",
            resultados
        )

    # --------------------------------------------------------
    # EMITENTE = SAÍDA
    # --------------------------------------------------------

    resultados = []

    if emit_cnpj:

        for cliente in clientes:

            if (
                cliente["cnpj"]
                == emit_cnpj
            ):

                resultados.append(
                    (
                        "SAIDAS",
                        cliente
                    )
                )

    if len(resultados) == 1:
        return resultados[0]

    if len(resultados) > 1:

        return (
            "AMBIGUO",
            resultados
        )

    return None


# ============================================================
# PASTA DE DESTINO
# ============================================================

def pasta_destino(
    cliente,
    data,
    direcao
):

    if not data:
        return None

    try:

        dia, mes, ano = data.split("/")

        ano_curto = ano[-2:]

        pasta = (
            cliente["caminho"]
            / "NF"
            / direcao
            / f"{mes}-{ano_curto}"
        )

        return pasta

    except Exception:
        return None


# ============================================================
# PDF — LEITURA
# ============================================================

def extrair_texto_pdf(caminho):

    textos = []

    try:

        reader = PdfReader(
            str(caminho)
        )

        for pagina in reader.pages:

            try:

                texto = (
                    pagina.extract_text()
                    or ""
                )

                if texto:
                    textos.append(texto)

            except Exception:
                pass

    except Exception:
        return ""

    return "\n".join(textos)


# ============================================================
# PDF — NÚMEROS
# ============================================================

def extrair_numeros_pdf(texto):
    """V10: várias grafias de rótulos fiscais; evita números sem contexto."""
    if not texto:
        return []

    texto_upper = re.sub(r"[\u00a0\t]+", " ", str(texto).upper())
    padroes = [
        # Nota fiscal / NFS-e
        r"\bN[ÚU]MERO\s+(?:DA\s+)?(?:NOTA|NOTA FISCAL|NFS[ -]?E|NFSE|DOCUMENTO)\s*(?:FISCAL)?\s*(?:N[º°O.]?\s*)?[:#\-]?\s*(\d{1,12})\b",
        r"\bN[º°O.]\s*(?:DA\s+)?(?:NOTA|NOTA FISCAL|NFS[ -]?E|NFSE|DOCUMENTO)\s*[:#\-]?\s*(\d{1,12})\b",
        r"\b(?:NOTA FISCAL DE SERVI[CÇ]OS?|NOTA FISCAL DE SERVI[CÇ]O ELETR[ÔO]NICA|DANFSE)\s*(?:N[º°O.]?\s*)?[:#\-]?\s*(\d{1,12})\b",
        r"\bNFS[ -]?E\s*(?:N[º°O.]?\s*)?[:#\-]?\s*(\d{1,12})\b",
        r"\bNFSE\s*(?:N[º°O.]?\s*)?[:#\-]?\s*(\d{1,12})\b",
        r"\bN[ÚU]MERO\s+DA\s+NOTA\s*\n\s*(\d{1,12})\b",
        # Alguns layouts colocam o rótulo depois do número.
        r"\b(\d{1,12})\s*(?:[-–:]\s*)?(?:N[º°O.]?\s*)?(?:DA\s+)?NOTA FISCAL\b",
    ]

    encontrados = []
    for padrao in padroes:
        for match in re.finditer(padrao, texto_upper):
            numero = limpar_numero(match.group(1))
            if numero and numero != "0" and numero not in encontrados:
                encontrados.append(numero)

    # Redundância segura: número da NF-e obtido de uma chave de acesso válida.
    for chave in encontrar_chaves(texto_upper):
        if validar_chave_nfe(chave):
            numero = numero_por_chave_nfe(chave)
            if numero and numero != "0" and numero not in encontrados:
                encontrados.append(numero)

    return encontrados

def extrair_cnpjs_pdf(texto):

    encontrados = set()

    for numero in re.findall(
        r"\d[\d./-]{12,20}\d",
        texto
    ):

        cnpj = limpar_cnpj(
            numero
        )

        if len(cnpj) == 14:

            encontrados.add(
                cnpj
            )

    return sorted(
        encontrados
    )


# ============================================================
# PDF — INDEXAÇÃO
# ============================================================

def indexar_pdfs():

    indice = []

    for caminho in PASTA_PROCESSAR.glob(
        "*.pdf"
    ):

        texto = extrair_texto_pdf(
            caminho
        )

        # ----------------------------------------------------
        # CHAVES
        # ----------------------------------------------------

        chaves = encontrar_chaves(
            texto
            + " "
            + caminho.stem
        )

        # ----------------------------------------------------
        # NÚMEROS
        # ----------------------------------------------------

        numeros = extrair_numeros_pdf(
            texto
        )

        # ----------------------------------------------------
        # CNPJs
        # ----------------------------------------------------

        cnpjs = extrair_cnpjs_pdf(
            texto
        )

        # ----------------------------------------------------
        # DATAS
        # ----------------------------------------------------

        datas = []

        for m in re.finditer(
            r"\b\d{2}/\d{2}/\d{4}\b",
            texto
        ):

            datas.append(
                m.group(0)
            )

        indice.append({

            "arquivo": caminho,

            "texto": texto,

            "chaves": chaves,

            "numeros": numeros,

            "cnpjs": cnpjs,

            "datas": sorted(
                set(datas)
            ),

            "usado": False
        })

    return indice


# ============================================================
# ASSOCIAR PDF
# ============================================================

def associar_pdf(
    xml_info,
    pdfs
):

    chave = xml_info["chave"]
    numero = xml_info["numero"]

    # --------------------------------------------------------
    # 1 — CHAVE EXATA
    # --------------------------------------------------------

    if chave:

        encontrados = []

        for pdf in pdfs:

            if pdf["usado"]:
                continue

            if chave in pdf["chaves"]:

                encontrados.append(
                    pdf
                )

        if len(encontrados) == 1:

            pdf = encontrados[0]

            pdf["usado"] = True

            return (
                pdf,
                "CHAVE_EXATA"
            )

        if len(encontrados) > 1:

            return (
                None,
                "CONFLITO_CHAVE"
            )

    # --------------------------------------------------------
    # 2 — NÚMERO EXATO
    # --------------------------------------------------------

    if numero:

        encontrados = []

        for pdf in pdfs:

            if pdf["usado"]:
                continue

            if numero in pdf["numeros"]:

                encontrados.append(
                    pdf
                )

        if len(encontrados) == 1:

            pdf = encontrados[0]

            pdf["usado"] = True

            return (
                pdf,
                "NUMERO_EXATO"
            )

        if len(encontrados) > 1:

            # Tenta resolver pelo CNPJ

            cnpjs_xml = {
                xml_info["emit_cnpj"],
                xml_info["dest_cnpj"]
            }

            cnpjs_xml.discard(
                None
            )

            filtrados = []

            for pdf in encontrados:

                if cnpjs_xml.intersection(
                    set(pdf["cnpjs"])
                ):

                    filtrados.append(
                        pdf
                    )

            if len(filtrados) == 1:

                pdf = filtrados[0]

                pdf["usado"] = True

                return (
                    pdf,
                    "NUMERO_E_CNPJ"
                )

            return (
                None,
                "CONFLITO_NUMERO"
            )

    # --------------------------------------------------------
    # 3 — CNPJ + DATA
    # --------------------------------------------------------

    cnpjs_xml = {
        xml_info["emit_cnpj"],
        xml_info["dest_cnpj"]
    }

    cnpjs_xml.discard(
        None
    )

    if (
        cnpjs_xml
        and xml_info["data"]
    ):

        candidatos = []

        for pdf in pdfs:

            if pdf["usado"]:
                continue

            if not cnpjs_xml.intersection(
                set(pdf["cnpjs"])
            ):
                continue

            if (
                xml_info["data"]
                not in pdf["datas"]
            ):
                continue

            candidatos.append(
                pdf
            )

        if len(candidatos) == 1:

            pdf = candidatos[0]

            pdf["usado"] = True

            return (
                pdf,
                "CNPJ_DATA"
            )

    return (
        None,
        "SEM_ASSOCIACAO"
    )


# ============================================================
# PROCESSAMENTO XML
# ============================================================

def processar_xml(
    caminho_xml,
    clientes,
    pdfs
):

    root = ler_xml(
        caminho_xml
    )

    if root is None:

        return {

            "arquivo":
                caminho_xml.name,

            "status":
                "REVISÃO",

            "motivo":
                "XML inválido ou ilegível"
        }

    # --------------------------------------------------------
    # TIPO
    # --------------------------------------------------------

    tipo = identificar_tipo_xml(
        root
    )

    if not tipo:

        return {

            "arquivo":
                caminho_xml.name,

            "status":
                "REVISÃO",

            "motivo":
                "Tipo de documento não identificado"
        }

    # --------------------------------------------------------
    # CHAVE
    # --------------------------------------------------------

    chave = extrair_chave_xml(
        root,
        caminho_xml
    )

    # --------------------------------------------------------
    # NÚMERO
    # --------------------------------------------------------

    (
        numero,
        candidatos_numero,
        metodo_numero
    ) = decidir_numero_xml(
        root,
        tipo,
        caminho_xml
    )

    # --------------------------------------------------------
    # DATA
    # --------------------------------------------------------

    data = extrair_data_xml(
        root
    )

    # --------------------------------------------------------
    # CNPJs
    # --------------------------------------------------------

    (
        emit_cnpj,
        dest_cnpj
    ) = extrair_emitente_destinatario(
        root
    )

    # --------------------------------------------------------
    # CLIENTE
    # --------------------------------------------------------

    resultado_cliente = identificar_cliente(
        clientes,
        emit_cnpj,
        dest_cnpj
    )

    if resultado_cliente is None:
        # V10: mesmo sem cliente cadastrado, tenta localizar o PDF correspondente.
        # Isso ajuda a separar "documento encontrado" de "cliente identificado".
        info_sem_cliente = {
            "chave": chave,
            "numero": numero,
            "data": data,
            "emit_cnpj": emit_cnpj,
            "dest_cnpj": dest_cnpj
        }
        pdf_sem_cliente, metodo_pdf_sem_cliente = associar_pdf(
            info_sem_cliente,
            pdfs
        )

        return {
            "arquivo": caminho_xml.name,
            "tipo": tipo,
            "chave": chave,
            "numero": numero,
            "data": data,
            "emit_cnpj": emit_cnpj,
            "dest_cnpj": dest_cnpj,
            "pdf": pdf_sem_cliente,
            "metodo_pdf": metodo_pdf_sem_cliente,
            "status": "REVISÃO",
            "motivo": "CNPJ do cliente não encontrado no banco.xlsx; associação de PDF não autoriza movimentação",
            "candidatos_numero": candidatos_numero
        }

    if resultado_cliente[0] == "AMBIGUO":

        return {

            "arquivo":
                caminho_xml.name,

            "tipo":
                tipo,

            "chave":
                chave,

            "numero":
                numero,

            "data":
                data,

            "emit_cnpj":
                emit_cnpj,

            "dest_cnpj":
                dest_cnpj,

            "status":
                "REVISÃO",

            "motivo":
                "CNPJ associado a mais de um cliente",

            "candidatos_numero":
                candidatos_numero
        }

    direcao, cliente = resultado_cliente

    # --------------------------------------------------------
    # NÚMERO NÃO ENCONTRADO
    # --------------------------------------------------------

    if not numero:

        return {

            "arquivo":
                caminho_xml.name,

            "tipo":
                tipo,

            "chave":
                chave,

            "numero":
                None,

            "data":
                data,

            "emit_cnpj":
                emit_cnpj,

            "dest_cnpj":
                dest_cnpj,

            "cliente":
                cliente,

            "direcao":
                direcao,

            "status":
                "REVISÃO",

            "motivo":
                metodo_numero,

            "candidatos_numero":
                candidatos_numero
        }

    # --------------------------------------------------------
    # DATA NÃO ENCONTRADA
    # --------------------------------------------------------

    if not data:

        return {

            "arquivo":
                caminho_xml.name,

            "tipo":
                tipo,

            "chave":
                chave,

            "numero":
                numero,

            "data":
                None,

            "emit_cnpj":
                emit_cnpj,

            "dest_cnpj":
                dest_cnpj,

            "cliente":
                cliente,

            "direcao":
                direcao,

            "status":
                "REVISÃO",

            "motivo":
                "Data de emissão não identificada",

            "candidatos_numero":
                candidatos_numero
        }

    # --------------------------------------------------------
    # ASSOCIAR PDF
    # --------------------------------------------------------

    info = {

        "chave":
            chave,

        "numero":
            numero,

        "data":
            data,

        "emit_cnpj":
            emit_cnpj,

        "dest_cnpj":
            dest_cnpj
    }

    pdf, metodo_pdf = associar_pdf(
        info,
        pdfs
    )

    # --------------------------------------------------------
    # DESTINO
    # --------------------------------------------------------

    pasta = pasta_destino(
        cliente,
        data,
        direcao
    )

    # --------------------------------------------------------
    # NOME
    # --------------------------------------------------------

    if tipo == "NF-e":

        nome_base = (
            f"NF {numero}"
        )

    else:

        nome_base = (
            f"NFS {numero}"
        )

    arquivos_finais = [
        f"{nome_base}.xml"
    ]

    if pdf:

        arquivos_finais.append(
            f"{nome_base}.pdf"
        )

    # --------------------------------------------------------
    # RESULTADO
    # --------------------------------------------------------

    return {

        "arquivo":
            caminho_xml.name,

        "tipo":
            tipo,

        "chave":
            chave,

        "numero":
            numero,

        "data":
            data,

        "emit_cnpj":
            emit_cnpj,

        "dest_cnpj":
            dest_cnpj,

        "cliente":
            cliente,

        "direcao":
            direcao,

        "pasta":
            pasta,

        "nome_base":
            nome_base,

        "arquivos_finais":
            arquivos_finais,

        "pdf":
            pdf,

        "metodo_pdf":
            metodo_pdf,

        "status":
            "OK",

        "metodo_numero":
            metodo_numero,

        "candidatos_numero":
            candidatos_numero
    }


# ============================================================
# RELATÓRIO
# ============================================================

def escrever_resultado(
    f,
    resultado
):

    f.write(
        "\n"
        + "=" * 90
        + "\n"
    )

    f.write(
        f"ARQUIVO XML: "
        f"{resultado.get('arquivo')}\n"
    )

    f.write(
        f"STATUS: "
        f"{resultado.get('status')}\n"
    )

    if resultado.get("tipo"):

        f.write(
            f"TIPO: "
            f"{resultado.get('tipo')}\n"
        )

    if resultado.get("numero"):

        f.write(
            f"NÚMERO: "
            f"{resultado.get('numero')}\n"
        )

    if resultado.get("data"):

        f.write(
            f"DATA: "
            f"{resultado.get('data')}\n"
        )

    if resultado.get("chave"):

        f.write(
            f"CHAVE: "
            f"{resultado.get('chave')}\n"
        )

    if resultado.get("emit_cnpj"):

        f.write(
            f"EMITENTE: "
            f"{resultado.get('emit_cnpj')}\n"
        )

    if resultado.get("dest_cnpj"):

        f.write(
            f"DESTINATÁRIO: "
            f"{resultado.get('dest_cnpj')}\n"
        )

    cliente = resultado.get(
        "cliente"
    )

    if isinstance(
        cliente,
        dict
    ):

        f.write(
            f"CLIENTE: "
            f"{cliente.get('cod')} - "
            f"{cliente.get('razao')}\n"
        )

        f.write(
            f"CAMINHO: "
            f"{cliente.get('caminho')}\n"
        )

    if resultado.get("direcao"):

        f.write(
            f"DIREÇÃO: "
            f"{resultado.get('direcao')}\n"
        )

    if resultado.get(
        "metodo_numero"
    ):

        f.write(
            f"MÉTODO NÚMERO: "
            f"{resultado.get('metodo_numero')}\n"
        )

    if resultado.get(
        "candidatos_numero"
    ):

        f.write(
            "CANDIDATOS DE NÚMERO:\n"
        )

        for candidato in (
            resultado[
                "candidatos_numero"
            ]
        ):

            f.write(
                f"  - "
                f"{candidato.get('numero')} "
                f"| "
                f"{candidato.get('metodo')} "
                f"| "
                f"peso="
                f"{candidato.get('peso')}\n"
            )

    if resultado.get("pdf"):

        pdf = resultado["pdf"]

        f.write(
            f"PDF ASSOCIADO: "
            f"{pdf['arquivo'].name}\n"
        )

        f.write(
            f"MÉTODO PDF: "
            f"{resultado.get('metodo_pdf')}\n"
        )

    elif resultado.get(
        "status"
    ) == "OK":

        f.write(
            "PDF ASSOCIADO: "
            "NÃO ENCONTRADO "
            "(XML pode existir sozinho)\n"
        )

    if resultado.get("pasta"):

        f.write(
            f"PASTA DESTINO: "
            f"{resultado.get('pasta')}\n"
        )

    if resultado.get(
        "nome_base"
    ):

        f.write(
            f"NOME PADRONIZADO: "
            f"{resultado.get('nome_base')}\n"
        )

    if resultado.get(
        "arquivos_finais"
    ):

        f.write(
            "ARQUIVOS FINAIS:\n"
        )

        for nome in resultado[
            "arquivos_finais"
        ]:

            f.write(
                f"  - {nome}\n"
            )

    if resultado.get(
        "motivo"
    ):

        f.write(
            f"MOTIVO: "
            f"{resultado.get('motivo')}\n"
        )


# ============================================================
# MAIN
# ============================================================

def main():

    print(
        "=============================================="
    )

    print(
        " PROCESSADOR XML + PDF "
        "NF-e / NFS-e - V10"
    )

    print(
        " MODO SIMULAÇÃO"
    )

    print(
        "=============================================="
    )

    # --------------------------------------------------------
    # VERIFICAÇÕES
    # --------------------------------------------------------

    if not PASTA_PROCESSAR.exists():

        print(
            f"Pasta não encontrada: "
            f"{PASTA_PROCESSAR}"
        )

        return

    if not ARQUIVO_BANCO.exists():

        print(
            f"Banco não encontrado: "
            f"{ARQUIVO_BANCO}"
        )

        return

    # --------------------------------------------------------
    # CLIENTES
    # --------------------------------------------------------

    print(
        "Carregando banco de clientes..."
    )

    clientes = carregar_clientes()

    print(
        f"Clientes carregados: "
        f"{len(clientes)}"
    )

    # --------------------------------------------------------
    # XML
    # --------------------------------------------------------

    xmls = sorted(
        PASTA_PROCESSAR.glob(
            "*.xml"
        )
    )

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    print(
        "Indexando PDFs..."
    )

    pdfs = indexar_pdfs()

    print(
        f"PDFs encontrados: "
        f"{len(pdfs)}"
    )

    # --------------------------------------------------------
    # PROCESSAMENTO
    # --------------------------------------------------------

    resultados = []

    for xml in xmls:

        print()
        print(
            f"Processando XML: "
            f"{xml.name}"
        )

        resultado = processar_xml(
            xml,
            clientes,
            pdfs
        )

        resultados.append(
            resultado
        )

    # --------------------------------------------------------
    # PDFS ÓRFÃOS
    # --------------------------------------------------------

    pdfs_orfaos = [

        pdf
        for pdf in pdfs

        if not pdf["usado"]
    ]

    # --------------------------------------------------------
    # ESTATÍSTICAS
    # --------------------------------------------------------

    total_xml = len(
        xmls
    )

    total_pdf = len(
        pdfs
    )

    ok_xml = sum(
        1
        for resultado in resultados
        if resultado.get(
            "status"
        ) == "OK"
    )

    revisao_xml = (
        total_xml
        - ok_xml
    )

    pares = sum(
        1
        for resultado in resultados
        if resultado.get(
            "pdf"
        ) is not None
    )

    # ========================================================
    # RELATÓRIO
    # ========================================================

    with open(
        ARQUIVO_RELATORIO,
        "w",
        encoding="utf-8"
    ) as f:

        f.write(
            "RELATÓRIO DE SIMULAÇÃO - V10\n"
        )

        f.write(
            "Nenhum arquivo foi "
            "movido ou renomeado.\n\n"
        )

        f.write(
            f"XMLs encontrados: "
            f"{total_xml}\n"
        )

        f.write(
            f"PDFs encontrados: "
            f"{total_pdf}\n"
        )

        f.write(
            f"XMLs OK: "
            f"{ok_xml}\n"
        )

        f.write(
            f"XMLs em revisão: "
            f"{revisao_xml}\n"
        )

        f.write(
            f"XML/PDF associados: "
            f"{pares}\n"
        )

        f.write(
            f"PDFs órfãos: "
            f"{len(pdfs_orfaos)}\n"
        )

        pdfs_sem_texto = sum(1 for pdf in pdfs if not pdf.get("texto", "").strip())
        f.write(
            f"PDFs sem texto extraível (possível OCR): {pdfs_sem_texto}\n"
        )

        # ----------------------------------------------------
        # RESULTADOS XML
        # ----------------------------------------------------

        for resultado in resultados:

            escrever_resultado(
                f,
                resultado
            )

        # ----------------------------------------------------
        # ÓRFÃOS
        # ----------------------------------------------------

        f.write(
            "\n\n"
            + "#"
            * 90
            + "\n"
        )

        f.write(
            "PDFs ORFÃOS / NÃO ASSOCIADOS\n"
        )

        f.write(
            "#"
            * 90
            + "\n"
        )

        for pdf in pdfs_orfaos:

            f.write(
                "\n"
            )

            f.write(
                f"PDF: "
                f"{pdf['arquivo'].name}\n"
            )

            f.write(
                f"CHAVES: {', '.join(pdf['chaves']) if pdf['chaves'] else 'nenhuma'}\n"
            )

            f.write(
                f"NÚMEROS: {', '.join(pdf['numeros']) if pdf['numeros'] else 'nenhum'}\n"
            )

            f.write(
                f"CNPJs: {', '.join(pdf['cnpjs']) if pdf['cnpjs'] else 'nenhum'}\n"
            )

            f.write(
                f"DATAS: {', '.join(pdf['datas']) if pdf['datas'] else 'nenhuma'}\n"
            )

            f.write(
                "DIAGNÓSTICO: Nenhum XML correspondente encontrado neste lote pelos critérios atuais; "
                "conferir se o XML foi baixado e se pertence à mesma competência.\n"
            )
            if not pdf.get("texto", "").strip():
                f.write(
                    "OBSERVAÇÃO: não foi possível extrair texto do PDF; pode ser digitalizado e exigir OCR.\n"
                )

        # ----------------------------------------------------
        # FINAL
        # ----------------------------------------------------

        f.write(
            "\n\n"
            + "#"
            * 90
            + "\n"
        )

        f.write(
            "FIM DO RELATÓRIO\n"
        )

        f.write(
            "MODO SIMULAÇÃO — "
            "NENHUM ARQUIVO FOI ALTERADO.\n"
        )

    # ========================================================
    # CONSOLE
    # ========================================================

    print()

    print(
        "=============================================="
    )

    print(
        " SIMULAÇÃO CONCLUÍDA"
    )

    print(
        "=============================================="
    )

    print(
        f"XMLs: "
        f"{total_xml}"
    )

    print(
        f"PDFs: "
        f"{total_pdf}"
    )

    print(
        f"XMLs OK: "
        f"{ok_xml}"
    )

    print(
        f"XMLs revisão: "
        f"{revisao_xml}"
    )

    print(
        f"Pares XML/PDF: "
        f"{pares}"
    )

    print(
        f"PDFs órfãos: "
        f"{len(pdfs_orfaos)}"
    )

    print()

    print(
        f"Relatório: "
        f"{ARQUIVO_RELATORIO}"
    )

    print()

    print(
        "NENHUM ARQUIVO FOI "
        "MOVIDO OU RENOMEADO."
    )


# ============================================================
# EXECUÇÃO
# ============================================================

if __name__ == "__main__":
    main()

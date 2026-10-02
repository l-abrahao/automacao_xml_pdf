# -*- coding: utf-8 -*-
"""
CRIAR ESTRUTURA DE PASTAS DOS CLIENTES

Lê a planilha "banco.xlsx" e utiliza a coluna "CAMINHO ABSOLUTO"
para criar as pastas que estiverem faltando.

Estrutura criada para 2026:
CLIENTE
├── NF
│   ├── ENTRADAS
│   │   └── 2026
│   │       ├── 01-26 ... 12-26
│   └── SAIDAS
│       └── 2026
│           ├── 01-26 ... 12-26
└── PGDAS
    └── 2026
        ├── 01-26 ... 12-26

O script NÃO apaga, move ou altera arquivos existentes.
Se a pasta principal do cliente não existir, ela será apenas informada
no relatório para evitar criar uma pasta no caminho errado.
"""

from pathlib import Path
import openpyxl
from datetime import datetime

# ============================================================
# CONFIGURAÇÃO
# ============================================================

ANO = 2026

# A planilha deve ficar na mesma pasta deste script.
ARQUIVO_EXCEL = Path(__file__).resolve().parent / "banco.xlsx"

# ============================================================
# LEITURA DA PLANILHA
# ============================================================

if not ARQUIVO_EXCEL.exists():
    print()
    print("ERRO: Não encontrei o arquivo:")
    print(ARQUIVO_EXCEL)
    print()
    print("Coloque o banco.xlsx na mesma pasta deste script.")
    input("Pressione ENTER para fechar...")
    raise SystemExit(1)

try:
    wb = openpyxl.load_workbook(ARQUIVO_EXCEL, data_only=True)
except Exception as e:
    print(f"ERRO ao abrir a planilha: {e}")
    input("Pressione ENTER para fechar...")
    raise SystemExit(1)

ws = wb[wb.sheetnames[0]]

# Localiza os cabeçalhos
cabecalhos = {}
for col in range(1, ws.max_column + 1):
    valor = ws.cell(1, col).value
    if valor is not None:
        cabecalhos[str(valor).strip().upper()] = col

col_caminho = cabecalhos.get("CAMINHO ABSOLUTO")

if not col_caminho:
    print("ERRO: A planilha não possui a coluna 'CAMINHO ABSOLUTO'.")
    print("Cabeçalhos encontrados:", list(cabecalhos.keys()))
    input("Pressione ENTER para fechar...")
    raise SystemExit(1)

# ============================================================
# CRIAÇÃO DAS PASTAS
# ============================================================

meses = [f"{mes:02d}-{str(ANO)[-2:]}" for mes in range(1, 13)]

criadas = []
ja_existiam = []
erros = []

for linha in range(2, ws.max_row + 1):
    codigo = ws.cell(linha, cabecalhos.get("COD", 1)).value
    razao = ws.cell(linha, cabecalhos.get("RAZÃO SOCIAL", 2)).value
    caminho_valor = ws.cell(linha, col_caminho).value

    if not caminho_valor:
        erros.append(
            f"Linha {linha}: caminho absoluto vazio "
            f"({razao or 'cliente sem nome'})."
        )
        continue

    caminho_cliente = Path(str(caminho_valor).strip())

    # Não cria a pasta principal automaticamente.
    # Isso evita criar uma pasta errada caso o caminho esteja digitado incorretamente.
    if not caminho_cliente.exists():
        erros.append(
            f"Linha {linha}: pasta do cliente NÃO encontrada: "
            f"{caminho_cliente}"
        )
        continue

    pastas = [
        caminho_cliente / "NF" / "ENTRADAS" / str(ANO),
        caminho_cliente / "NF" / "SAIDAS" / str(ANO),
        caminho_cliente / "PGDAS" / str(ANO),
    ]

    for pasta_base in pastas:
        # Cria a pasta base, se necessário
        if not pasta_base.exists():
            try:
                pasta_base.mkdir(parents=True, exist_ok=True)
                criadas.append(str(pasta_base))
            except Exception as e:
                erros.append(f"Erro ao criar {pasta_base}: {e}")
                continue

        # Cria os 12 meses
        for mes in meses:
            pasta_mes = pasta_base / mes

            if pasta_mes.exists():
                ja_existiam.append(str(pasta_mes))
            else:
                try:
                    pasta_mes.mkdir(parents=True, exist_ok=True)
                    criadas.append(str(pasta_mes))
                except Exception as e:
                    erros.append(f"Erro ao criar {pasta_mes}: {e}")

# ============================================================
# RELATÓRIO
# ============================================================

relatorio = ARQUIVO_EXCEL.parent / "relatorio_criacao_pastas.txt"

with open(relatorio, "w", encoding="utf-8") as arquivo:
    arquivo.write("RELATÓRIO - CRIAÇÃO DE PASTAS\n")
    arquivo.write("=" * 70 + "\n")
    arquivo.write(f"Data/hora: {datetime.now():%d/%m/%Y %H:%M:%S}\n")
    arquivo.write(f"Ano processado: {ANO}\n")
    arquivo.write(f"Planilha: {ARQUIVO_EXCEL}\n\n")

    arquivo.write(f"PASTAS CRIADAS: {len(criadas)}\n")
    arquivo.write("-" * 70 + "\n")
    for item in criadas:
        arquivo.write(item + "\n")

    arquivo.write(f"\nPASTAS QUE JÁ EXISTIAM: {len(ja_existiam)}\n")
    arquivo.write("-" * 70 + "\n")
    for item in ja_existiam:
        arquivo.write(item + "\n")

    arquivo.write(f"\nERROS/PENDÊNCIAS: {len(erros)}\n")
    arquivo.write("-" * 70 + "\n")
    for item in erros:
        arquivo.write(item + "\n")

print()
print("=" * 70)
print("  ESTRUTURA DE PASTAS PROCESSADA")
print("=" * 70)
print(f"Ano: {ANO}")
print(f"Pastas criadas:       {len(criadas)}")
print(f"Pastas já existentes: {len(ja_existiam)}")
print(f"Erros/pendências:     {len(erros)}")
print()
print(f"Relatório salvo em:")
print(relatorio)
print()

if erros:
    print("ATENÇÃO: existem pendências. Consulte o relatório.")
else:
    print("Tudo processado sem erros.")

input("\nPressione ENTER para fechar...")
